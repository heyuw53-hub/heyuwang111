"""Nettoyage des données, statistiques, histogrammes + gaussienne, export Excel.

Reprend la logique de Nettoyage_Polaro.py / Nettoyage_UV.py / Distribution_Normale.py.
"""
import io

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from config import N_MIN_GAUSS



def valeurs_aberrantes(df: pd.DataFrame, cols: list[str], seuil: float = 0.30,
                       reference: str = "moyenne") -> pd.DataFrame:
    """Tableau booléen de même forme que df[cols] : True = valeur hors de ±seuil de la référence.

    reference = "moyenne" (comme les scripts d'origine) ou "médiane"
    (plus robuste : une valeur aberrante ne décale pas la fenêtre).
    Le nettoyage ne s'applique qu'à partir de N_MIN_GAUSS valeurs.
    """
    masque = pd.DataFrame(False, index=df.index, columns=cols)
    for c in cols:
        x = df[c]
        if x.notna().sum() < N_MIN_GAUSS:
            continue  # trop peu de valeurs : la moyenne n'est pas une référence fiable, on ne marque rien
        ref = x.mean() if reference == "moyenne" else x.median()
        bas, haut = sorted((ref * (1 - seuil), ref * (1 + seuil)))
        masque[c] = x.notna() & ~x.between(bas, haut)
    return masque


def nettoyer(df: pd.DataFrame, cols: list[str], seuil: float = 0.30, reference: str = "moyenne"):
    """Garde les lignes dont TOUTES les valeurs sont à ±seuil de la référence.
    Retourne (df_gardé, df_exclu).
    """
    d = df.dropna(subset=cols)
    exclu = valeurs_aberrantes(d, cols, seuil, reference).any(axis=1)
    return d[~exclu], d[exclu]


# ---------------------------------------------------------------------------
# Incertitudes (document ECPM-TPSA-MOP-003 « Traitement des données »)
# ---------------------------------------------------------------------------
# Coefficient de Student à 95 % en fonction du nombre de degrés de liberté.
# 1 à 14 : tableau du document ; au-delà : valeurs usuelles des tables de Student.
STUDENT_95 = {1: 12.7, 2: 4.30, 3: 3.18, 4: 2.78, 5: 2.57, 6: 2.45, 7: 2.36, 8: 2.31, 9: 2.26,
              10: 2.23, 11: 2.20, 12: 2.18, 13: 2.16, 14: 2.14, 15: 2.13, 16: 2.12, 17: 2.11,
              18: 2.10, 19: 2.09, 20: 2.09, 25: 2.06, 30: 2.04, 40: 2.02, 60: 2.00, 120: 1.98}


def t_student_95(ddl: int) -> float:
    """t à 95 % ; entre deux valeurs tabulées on prend la plus grande (prudent)."""
    if ddl < 1:
        return float("nan")
    tabules = [k for k in STUDENT_95 if k <= ddl]
    return STUDENT_95[max(tabules)] if ddl <= 120 else 1.96


def arrondi_resultat(valeur: float, incertitude: float) -> str:
    """Règles du document : incertitude à 1 chiffre significatif arrondie PAR EXCÈS,
    résultat arrondi au même rang décimal. Ex. (45.123456, 0.018453) -> '45.12 ± 0.02'."""
    if not (np.isfinite(valeur) and np.isfinite(incertitude)) or incertitude <= 0:
        return "–"
    e = int(np.floor(np.log10(incertitude)))
    inc = np.ceil(round(incertitude / 10 ** e, 9)) * 10 ** e   # round() évite 2.0000001 -> 3
    e = int(np.floor(np.log10(inc)))                            # 9.2 -> 10 change de rang
    dec = max(0, -e)
    return f"{round(valeur, -e):.{dec}f} ± {inc:.{dec}f}"


COLONNES_NUMERIQUES = ["Valeur théorique", "Écart relatif (%)", "Moyenne", "Écart-type", "CV (%)", "Min", "Max"]


def _stats(x: pd.Series, theo: float | None) -> dict:
    """n, moyenne, écart-type, CV (+ écart relatif à la valeur théorique). Pas d'incertitude
    ici : les incertitudes sont calculées par binôme, à la saisie (incertitudes.py)."""
    x = x.dropna()
    moy = x.mean() if len(x) else np.nan
    ecart = x.std(ddof=1) if len(x) > 1 else np.nan
    r = {"n": len(x), "Moyenne": moy, "Écart-type": ecart, "CV (%)": 100 * ecart / moy if moy else np.nan}
    if theo:
        r["Valeur théorique"] = theo
        r["Écart relatif (%)"] = 100 * (moy - theo) / theo
    return r


def statistiques(df: pd.DataFrame, champs: list[dict]) -> pd.DataFrame:
    lignes = []
    for c in champs:
        x = df[c["col"]].dropna()
        lignes.append({"Grandeur": c["label"], **_stats(x, c.get("theorique")),
                       "Min": x.min() if len(x) else np.nan, "Max": x.max() if len(x) else np.nan})
    return pd.DataFrame(lignes)


# ---------------------------------------------------------------------------
# Graphiques
# ---------------------------------------------------------------------------
# Palette catégorielle (ordre fixe ; la couleur suit l'entité, jamais son rang)
PALETTE = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300"]
BARRES = "#9ec5f4"      # histogramme (bleu clair)
COURBE = "#1c5cab"      # gaussienne (bleu foncé)
ENCRE = "#3d3d3a"       # texte / médianes
ROUGE = "#d62728"       # point de l'étudiant

plt.rcParams.update({
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.edgecolor": "#b5b4ad", "axes.labelcolor": ENCRE, "xtick.color": ENCRE, "ytick.color": ENCRE,
    "axes.grid": True, "grid.color": "#e8e7e1", "grid.linewidth": 0.8, "axes.axisbelow": True,
    "font.size": 9, "axes.titlesize": 10, "legend.frameon": False,
})


def _axes_lisibles(fig):
    """Notation scientifique pour les très petites/grandes valeurs, 5 graduations max."""
    from matplotlib.ticker import MaxNLocator, ScalarFormatter
    for ax in fig.axes:
        for axis in (ax.xaxis, ax.yaxis):
            if isinstance(axis.get_major_formatter(), ScalarFormatter):
                f = ScalarFormatter(useMathText=True)
                f.set_powerlimits((-3, 4))
                axis.set_major_formatter(f)
                axis.set_major_locator(MaxNLocator(5))


def couleur(i: int) -> str:
    return PALETTE[i % len(PALETTE)]


def _mu_sigma(x: np.ndarray, n_min: int = N_MIN_GAUSS):
    """(μ, σ) si assez de données pour une gaussienne, sinon None."""
    if len(x) >= max(2, n_min) and np.std(x, ddof=1) > 0:
        return float(np.mean(x)), float(np.std(x, ddof=1))
    return None


def _pdf(xs, mu, sigma):
    return np.exp(-0.5 * ((xs - mu) / sigma) ** 2) / (sigma * np.sqrt(2 * np.pi))


def _xs(valeurs: list[np.ndarray], extra: list[float] = ()):
    """Grille commune couvrant toutes les séries (et d'éventuels points)."""
    bornes = []
    for x in valeurs:
        ms = _mu_sigma(x)
        if len(x):
            bornes += [x.min(), x.max()]
        if ms:
            bornes += [ms[0] - 4 * ms[1], ms[0] + 4 * ms[1]]
    bornes += list(extra)
    if not bornes:
        return None
    lo, hi = min(bornes), max(bornes)
    if lo == hi:
        lo, hi = lo - 1, hi + 1
    return np.linspace(lo, hi, 400)


def _vide(ax, texte="Aucune donnée"):
    ax.text(0.5, 0.5, texte, ha="center", va="center", transform=ax.transAxes, color=ENCRE)


THEO = "#7b3fb5"        # valeur théorique (violet : hors palette des groupes/années)


def _ligne_theorique(ax, theo: float | None):
    if theo:
        ax.axvline(theo, color=THEO, lw=2, ls="-.", label=f"Valeur théorique ({theo:g})", zorder=4)
        ax.legend(loc="best")


def _histo_gauss(ax, x: np.ndarray, label: str, point: float | None = None, theo: float | None = None):
    """Histogramme + gaussienne ; `point` = valeur d'un étudiant (point rouge sur la courbe)."""
    ax.set_xlabel(label)
    ax.set_ylabel("Densité")
    if len(x) == 0:
        _vide(ax)
        return
    ax.hist(x, bins="auto", density=True, color=BARRES, edgecolor="white", linewidth=1.5)
    _ligne_theorique(ax, theo)
    ms = _mu_sigma(x)
    if not ms:
        ax.set_title(f"n = {len(x)} : gaussienne affichée à partir de {N_MIN_GAUSS} binômes")
        if point is not None:  # le binôme se voit quand même, sur l'axe
            ax.plot([point], [0], "o", ms=10, color=ROUGE, mec="white", mew=2, zorder=5,
                    clip_on=False, label="Votre binôme")
            ax.axvline(point, color=ROUGE, lw=1, alpha=0.5)
            ax.legend(loc="upper right")
        return
    mu, sigma = ms
    xs = _xs([x], [v for v in (point, theo) if v is not None])
    ax.plot(xs, _pdf(xs, mu, sigma), color=COURBE, lw=2)
    ax.axvline(mu, color=COURBE, ls="--", lw=1)
    ax.set_title(f"μ = {mu:.5g}   σ = {sigma:.3g}   n = {len(x)}")
    if point is not None:
        y = _pdf(point, mu, sigma)
        ax.plot([point], [y], "o", ms=10, color=ROUGE, mec="white", mew=2, zorder=5, label="Votre binôme")
        ax.axvline(point, color=ROUGE, lw=1, alpha=0.5)
        ax.legend(loc="best")


# matplotlib n'est pas thread-safe (analyseur mathtext notamment) ; Streamlit exécute
# plusieurs sessions en parallèle. On dessine ET on rend l'image sous un verrou global.
import threading as _threading

_VERROU = _threading.Lock()


def rendre_png(fonction, *args, **kwargs) -> bytes:
    """Appelle une fonction figure_*(...) et renvoie l'image PNG, de façon thread-safe."""
    with _VERROU:
        fig = fonction(*args, **kwargs)
        try:
            buf = io.BytesIO()
            fig.savefig(buf, format="png", dpi=110, bbox_inches="tight")
            return buf.getvalue()
        finally:
            plt.close(fig)


def figure_distribution(df: pd.DataFrame, champs: list[dict], point: dict | None = None):
    """Un histogramme (densité) par grandeur + loi normale ajustée (μ, σ des données).
    `point` : {col: valeur} → la valeur est marquée d'un point rouge sur la courbe."""
    n = len(champs)
    fig, axes = plt.subplots(1, n, figsize=(5 * n, 4), squeeze=False)
    for ax, c in zip(axes[0], champs):
        x = df[c["col"]].dropna().to_numpy(dtype=float)
        _histo_gauss(ax, x, c["label"], None if point is None else point.get(c["col"]), c.get("theorique"))
    _axes_lisibles(fig)
    fig.tight_layout()
    return fig


# matplotlib ≥ 3.10 : orientation= ; versions antérieures : vert=
import inspect as _inspect
_HORIZONTAL = ({"orientation": "horizontal"} if "orientation" in _inspect.signature(plt.Axes.boxplot).parameters
               else {"vert": False})


def _jitter(n, largeur=0.12, graine=0):
    return np.random.default_rng(graine).uniform(-largeur, largeur, n)


def figure_intra_groupe(df: pd.DataFrame, champs: list[dict]):
    """Un groupe : gaussienne (haut) + boîte à moustaches horizontale sur le même axe x (bas)."""
    n = len(champs)
    fig, axes = plt.subplots(2, n, figsize=(5 * n, 5.5), squeeze=False, sharex="col",
                             gridspec_kw={"height_ratios": [3, 1.3]})
    for j, c in enumerate(champs):
        x = df[c["col"]].dropna().to_numpy(dtype=float)
        _histo_gauss(axes[0, j], x, "", None, c.get("theorique"))
        if c.get("theorique"):
            axes[1, j].axvline(c["theorique"], color=THEO, lw=2, ls="-.")
        axes[0, j].set_xlabel("")
        ax = axes[1, j]
        ax.set_xlabel(c["label"])
        ax.grid(axis="y", visible=False)
        if len(x) == 0:
            ax.set_yticks([])
            continue
        ax.boxplot(x, **_HORIZONTAL, widths=0.5, patch_artist=True, showfliers=False,
                   boxprops=dict(facecolor=BARRES, edgecolor=COURBE), medianprops=dict(color=ENCRE, lw=2),
                   whiskerprops=dict(color=COURBE), capprops=dict(color=COURBE))
        ax.set_yticks([])
        ax.scatter(x, 1 + _jitter(len(x)), s=22, color=COURBE, edgecolor="white", linewidth=0.8, zorder=3)
        # étiquette = n° de binôme, pour repérer qui est où
        for v, yy, b in zip(x, 1 + _jitter(len(x)), df.loc[df[c["col"]].notna(), "binome"]):
            ax.annotate(str(b), (v, yy), xytext=(0, 6), textcoords="offset points",
                        ha="center", fontsize=7, color=ENCRE)
    _axes_lisibles(fig)
    fig.tight_layout()
    return fig


def figure_comparaison(df: pd.DataFrame, champs: list[dict], par: str, modalites: list, style: str):
    """Comparaison entre modalités (groupes ou années).
    Haut : une gaussienne par modalité, superposées.
    Bas  : style="boite" → une boîte à moustaches par modalité, côte à côte ;
           style="nuage" → nuage de points (x = modalité), avec moyenne ± σ."""
    n = len(champs)
    fig, axes = plt.subplots(2, n, figsize=(5.5 * n, 7.5), squeeze=False,
                             gridspec_kw={"height_ratios": [1, 1.1]})
    for j, c in enumerate(champs):
        series = [df.loc[df[par] == m, c["col"]].dropna().to_numpy(dtype=float) for m in modalites]
        # --- gaussiennes superposées
        ax = axes[0, j]
        ax.set_xlabel(c["label"])
        ax.set_ylabel("Densité")
        theo = c.get("theorique")
        xs = _xs(series, [theo] if theo else [])
        if xs is None:
            _vide(ax)
        else:
            for i, (m, x) in enumerate(zip(modalites, series)):
                ms = _mu_sigma(x)
                if ms:
                    ax.plot(xs, _pdf(xs, *ms), color=couleur(i), lw=2, label=f"{m} (n={len(x)})")
                    ax.axvline(ms[0], color=couleur(i), ls="--", lw=1)
                elif len(x):
                    ax.plot([], [], color=couleur(i), lw=2, label=f"{m} (n={len(x)}, < {N_MIN_GAUSS} : pas de courbe)")
            if theo:
                ax.axvline(theo, color=THEO, lw=2, ls="-.", label=f"Valeur théorique ({theo:g})")
            ax.legend(loc="upper left", bbox_to_anchor=(1.0, 1.0), fontsize=8)
        # --- boîtes ou nuage
        ax = axes[1, j]
        ax.set_ylabel(c["label"])
        pos = np.arange(1, len(modalites) + 1)
        if style == "boite":
            pleines = [(p, x, i) for i, (p, x) in enumerate(zip(pos, series)) if len(x)]
            if pleines:
                bp = ax.boxplot([x for _, x, _ in pleines], positions=[p for p, _, _ in pleines], widths=0.55,
                                patch_artist=True, showfliers=False, medianprops=dict(color=ENCRE, lw=2))
                for patch, (_, _, i) in zip(bp["boxes"], pleines):
                    patch.set(facecolor=couleur(i), alpha=0.35, edgecolor=couleur(i))
                for k, (p, x, i) in enumerate(pleines):
                    for part in ("whiskers", "caps"):
                        for line in bp[part][2 * k:2 * k + 2]:
                            line.set_color(couleur(i))
                    ax.scatter(p + _jitter(len(x), graine=k), x, s=18, color=couleur(i),
                               edgecolor="white", linewidth=0.8, zorder=3)
        else:
            for i, (p, x) in enumerate(zip(pos, series)):
                if not len(x):
                    continue
                ax.scatter(p + _jitter(len(x), 0.15, graine=i), x, s=22, color=couleur(i),
                           edgecolor="white", linewidth=0.8, zorder=3)
                ms = _mu_sigma(x, n_min=2)
                if ms:
                    ax.errorbar(p + 0.3, ms[0], yerr=ms[1], fmt="D", color=ENCRE, ms=6, capsize=4, lw=1.5, zorder=4)
            ax.plot([], [], "D", color=ENCRE, label="moyenne ± σ")
            ax.legend(loc="best", fontsize=8)
        if theo:
            ax.axhline(theo, color=THEO, lw=2, ls="-.", zorder=2)
        ax.set_xticks(pos, [str(m) for m in modalites])
        ax.set_xlim(0.4, len(modalites) + 0.8)
        ax.grid(axis="x", visible=False)
    _axes_lisibles(fig)
    fig.tight_layout()
    return fig


def statistiques_par(df: pd.DataFrame, champs: list[dict], par: str, modalites: list, nom: str) -> pd.DataFrame:
    """Même tableau que statistiques(), par grandeur et par modalité (groupe ou année)."""
    return pd.DataFrame([{"Grandeur": c["label"], nom: str(m), **_stats(df.loc[df[par] == m, c["col"]], c.get("theorique"))}
                         for c in champs for m in modalites])


def vers_format_modele(df: pd.DataFrame, champs: list[dict], incertitudes: list[dict] | None = None) -> pd.DataFrame:
    """Colonnes renommées pour l'affichage et l'export Excel.
    `incertitudes` (export Excel uniquement) : ajoute une colonne « a ± v » arrondie par binôme."""
    renommage = {"annee": "Année", "groupe": "Groupe", "binome": "Binôme",
                 **{c["col"]: c["label"] for c in champs}}
    vue = df[list(renommage)].rename(columns=renommage)
    for i in incertitudes or []:
        if i["col"] in df.columns:
            vue[i["label"]] = [arrondi_resultat(a, v) if pd.notna(v) else "–"
                               for a, v in zip(df[i["de"]], df[i["col"]])]
    return vue


JAUNE = "FFFF00"


def excel(df: pd.DataFrame, champs: list[dict], feuille: str, masque: pd.DataFrame | None = None,
          incertitudes: list[dict] | None = None) -> bytes:
    """Classeur à une seule feuille. Si `masque` est fourni (cf. valeurs_aberrantes),
    les cellules aberrantes sont surlignées en jaune ; aucune ligne n'est supprimée."""
    from openpyxl.styles import PatternFill

    vue = vers_format_modele(df, champs, incertitudes)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        vue.to_excel(w, sheet_name=feuille, index=False)
        ws = w.sheets[feuille]
        for j, nom in enumerate(vue.columns, start=1):  # largeur de colonne lisible
            ws.column_dimensions[ws.cell(1, j).column_letter].width = max(10, len(str(nom)) + 2)
        if masque is not None:
            jaune = PatternFill(start_color=JAUNE, end_color=JAUNE, fill_type="solid")
            for c in champs:
                j = vue.columns.get_loc(c["label"]) + 1
                for i, aberrant in enumerate(masque.loc[df.index, c["col"]].to_numpy(), start=2):
                    if aberrant:
                        ws.cell(i, j).fill = jaune
    return buf.getvalue()
